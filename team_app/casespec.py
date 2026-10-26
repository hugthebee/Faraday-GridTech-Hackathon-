"""Developer case: read the filled HACKATHON developer input template and apply it IN MEMORY to the packet code.

The packet's data files are never touched (the supplied checker hashes them). activate(spec) swaps the packet tables the evaluator and the
checker code read (resource options, case, response parameters, ramp, grid schedule, reference load profile) for copies carrying the
developer's values, and activate(None) restores the benchmark. Results under a non-default case are CUSTOM-CASE results: the checker's own
physical rules and cost accounting still verify them, but its input-consistency check against the released data is relaxed.

Applied: opening month and grid delay (app level), land (total, engine, fuel cell), relocation notice, interruption price, remote-site terms
(capacity, energy multiplier, tariff, service charge), nuclear PPA (month, price, max MW), mature peak demand (scales the reference load and
its tiers), load ramp, grid connection schedule.
NOT applied (needs the packet's tier-generation formula, which is not shipped): tier base shares, tier daily/weekly amplitudes, protected floor.
"""
import copy
from contextlib import contextmanager
import hashlib
import json
import sys
import threading
from dataclasses import dataclass, field, asdict
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import alt_adapter as A  # noqa: E402
import engine as E  # noqa: E402

gm = A.gm
cs = E.cs

# ----------------------------------------------------------------------------------------------- benchmark reference values
_R = cs.response
_O = cs.options.set_index('option_id')
REFERENCE = {
    'default_initial_service_month': int(cs.case['default_initial_service_month']),
    'grid_delay_months': 0,
    'onsite_energy_land_acres': float(cs.case['onsite_energy_land_acres']),
    'gas_engine_land_acres_per_mw': float(_O.loc['gas_engine', 'land_acres_per_mw']),
    'fuel_cell_land_acres_per_mw': float(_O.loc['fuel_cell', 'land_acres_per_mw']),
    'relocation_notice_minutes': float(_R['relocation_notice_minutes']),
    'interruption_consequence_usd_per_mwh': float(_R['interruption_consequence_usd_per_mwh']),
    'remote_capacity_mw': float(_R['remote_capacity_mw']),
    'remote_energy_multiplier': float(_R['remote_energy_multiplier']),
    'remote_electricity_usd_per_mwh': float(_R['remote_electricity_usd_per_mwh']),
    'relocation_service_usd_per_local_mwh': float(_R['relocation_service_usd_per_local_mwh']),
    'nuclear_available_month': int(_O.loc['nuclear', 'available_month']),
    'nuclear_payment_usd_per_mwh': float(_O.loc['nuclear', 'payment_usd_per_mwh']),
    'nuclear_max_mw': float(_O.loc['nuclear', 'max_mw']),
    'mature_peak_demand_mw': float(cs.hourly.load_mw.max()),
}
def _norm_events(ev):
    return sorted((pd.Timestamp(e['start_utc']).strftime('%Y-%m-%dT%H:%M'), pd.Timestamp(e['end_utc']).strftime('%Y-%m-%dT%H:%M'), round(float(e['import_fraction']), 6), round(float(e.get('notice_minutes', 30)), 3)) for e in ev)


_ev = pd.read_csv(HERE.parent / 'data' / 'stress_events.csv')
REF_EVENTS = [dict(id=str(r.event_id), start_utc=str(r.start_utc), end_utc=str(r.end_utc), import_fraction=float(r.import_fraction), notice_minutes=float(cs.hourly.notice_minutes.max()))
              for r in _ev.itertuples()]
REF_RAMP = [(int(m), float(f)) for m, f in zip(cs.ramp.months_since_initial_service, cs.ramp.load_fraction)]
REF_GRID = [(int(m), float(x)) for m, x in zip(cs.grid_schedule.from_month, cs.grid_schedule.max_import_mw)]
SCALAR_FIELDS = [k for k in REFERENCE if k not in ('default_initial_service_month', 'grid_delay_months')]   # these two are app-level, not model overrides
IGNORED_FIELDS = ['protected_infrastructure_floor_mw', 'tier1_base_share', 'tier1_daily_amplitude', 'tier1_weekly_amplitude',
                  'tier2_base_share', 'tier2_daily_amplitude', 'tier2_weekly_amplitude']
NUCLEAR_BLOCK = float(_O.loc['nuclear', 'block_mw'])


@dataclass
class CaseSpec:
    default_initial_service_month: int | None = None
    grid_delay_months: int | None = None
    onsite_energy_land_acres: float | None = None
    gas_engine_land_acres_per_mw: float | None = None
    fuel_cell_land_acres_per_mw: float | None = None
    relocation_notice_minutes: float | None = None
    interruption_consequence_usd_per_mwh: float | None = None
    remote_capacity_mw: float | None = None
    remote_energy_multiplier: float | None = None
    remote_electricity_usd_per_mwh: float | None = None
    relocation_service_usd_per_local_mwh: float | None = None
    nuclear_available_month: int | None = None
    nuclear_payment_usd_per_mwh: float | None = None
    nuclear_max_mw: float | None = None
    mature_peak_demand_mw: float | None = None
    ramp: list | None = None            # [(months_since_initial_service, load_fraction), ...]
    grid_schedule: list | None = None   # [(from_month, max_import_mw), ...]
    stress_events: list | None = None   # [{id, start_utc, end_utc, import_fraction, notice_minutes}, ...] replaces the packet's seven ERCOT events
    stress_area: str | None = None      # label only (which interconnection the events belong to)
    ignored: dict = field(default_factory=dict)
    warnings: list = field(default_factory=list)
    errors: list = field(default_factory=list)

    # ------------------------------------------------------------------------------------------ comparison with the benchmark
    def _val(self, k):
        v = getattr(self, k)
        return REFERENCE[k] if v is None else v

    def overrides(self):
        """Only the entries that actually differ from the packet benchmark (what changes the model)."""
        out = {}
        for k in SCALAR_FIELDS:
            v = getattr(self, k)
            if v is not None and abs(float(v) - REFERENCE[k]) > 1e-9:
                out[k] = float(v)
        if self.ramp is not None and [(int(m), round(float(f), 9)) for m, f in self.ramp] != [(m, round(f, 9)) for m, f in REF_RAMP]:
            out['ramp'] = [(int(m), float(f)) for m, f in self.ramp]
        if self.grid_schedule is not None and [(int(m), float(x)) for m, x in self.grid_schedule] != REF_GRID:
            out['grid_schedule'] = [(int(m), float(x)) for m, x in self.grid_schedule]
        if self.stress_events is not None and _norm_events(self.stress_events) != _norm_events(REF_EVENTS):
            out['stress_events'] = list(self.stress_events)
        return out

    def is_default(self):
        return not self.overrides()

    def key(self):
        return 'bench' if self.is_default() else hashlib.sha1(json.dumps(self.overrides(), sort_keys=True).encode()).hexdigest()[:10]

    def ramp_months(self):
        """Months from opening until demand reaches 100% (10 in the packet)."""
        ramp = self.ramp if self.ramp is not None else REF_RAMP
        for m, f in ramp:
            if f >= 1 - 1e-9:
                return int(m)
        return int(ramp[-1][0])

    def planned_opening(self):
        return self.default_initial_service_month if self.default_initial_service_month is not None else REFERENCE['default_initial_service_month']

    def to_dict(self):
        return asdict(self)

    @staticmethod
    def from_dict(d):
        return CaseSpec(**d)

    def diff_rows(self):
        rows = []
        for k in list(REFERENCE) + ['ramp', 'grid_schedule']:
            if k in ('ramp', 'grid_schedule'):
                v = getattr(self, k)
                ref = REF_RAMP if k == 'ramp' else REF_GRID
                rows.append(dict(parameter=k, packet='%d rows' % len(ref), yours=('(blank)' if v is None else '%d rows' % len(v)),
                                 status='changed' if k in self.overrides() else ('same' if v is not None else 'default'), note=''))
                continue
            v = getattr(self, k)
            app_level = k in ('default_initial_service_month', 'grid_delay_months')
            rows.append(dict(parameter=k, packet=REFERENCE[k], yours='(blank)' if v is None else v,
                             status='default' if v is None else ('changed' if (k in self.overrides() or (app_level and abs(float(v) - REFERENCE[k]) > 1e-9)) else 'same'),
                             note='sets the app default (not a model override)' if app_level else ''))
        for k, v in self.ignored.items():
            rows.append(dict(parameter=k, packet='-', yours=v, status='NOT APPLIED', note='needs the packet tier formula (not shipped)'))
        return rows


# ----------------------------------------------------------------------------------------------- template reader
def _num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    try:
        return float(str(v).replace(',', '').replace('$', '').strip())
    except ValueError:
        raise ValueError(f'not a number: {v!r}')


def from_template(src):
    """Parse the filled developer template (path or file-like). Blank cells = not provided = packet default."""
    import openpyxl
    wb = openpyxl.load_workbook(src, data_only=True)
    spec = CaseSpec()
    need = ['Project Inputs', 'Load Ramp', 'Grid Connection']
    missing = [s for s in need if s not in wb.sheetnames]
    if missing:
        spec.errors.append('Sheets not found: ' + ', '.join(missing))
        return spec
    ws = wb['Project Inputs']
    for row in ws.iter_rows(min_row=2):
        name = row[0].value
        if not isinstance(name, str) or name.startswith(tuple('12345')):
            continue
        name = name.strip()
        try:
            v = _num(row[3].value)
        except ValueError as exc:
            spec.errors.append(f'{name}: {exc}'); continue
        if v is None:
            continue
        if name in IGNORED_FIELDS:
            spec.ignored[name] = v
        elif name in REFERENCE:
            setattr(spec, name, v)
        else:
            spec.warnings.append(f'Unknown parameter {name!r} ignored')
    for k in ('default_initial_service_month', 'grid_delay_months', 'nuclear_available_month'):
        v = getattr(spec, k)
        if v is not None:
            if abs(v - round(v)) > 1e-9:
                spec.errors.append(f'{k} must be a whole number of months')
            setattr(spec, k, int(round(v)))
    # load ramp
    wl = wb['Load Ramp']
    try:
        peak = _num(wl['B3'].value)
        if peak is not None:
            spec.mature_peak_demand_mw = peak
    except ValueError as exc:
        spec.errors.append(f'mature_peak_demand_mw: {exc}')
    ramp, gap = [], False
    for r in range(6, wl.max_row + 1):
        m, f = wl.cell(r, 1).value, wl.cell(r, 2).value
        if not isinstance(m, (int, float)):
            continue
        try:
            f = _num(f)
        except ValueError as exc:
            spec.errors.append(f'ramp month {int(m)}: {exc}'); continue
        if f is None:
            gap = True
        elif gap:
            spec.errors.append('Load ramp has a blank row before month %d: fill rows from month 0 without gaps' % int(m))
        else:
            ramp.append((int(m), f))
    if ramp:
        spec.ramp = ramp
    # grid schedule
    wg = wb['Grid Connection']
    grid = []
    for r in range(5, wg.max_row + 1):
        a, b = wg.cell(r, 1).value, wg.cell(r, 2).value
        if a is None and b is None:
            continue
        try:
            a, b = _num(a), _num(b)
        except ValueError as exc:
            spec.errors.append(f'grid row {r}: {exc}'); continue
        if a is None or b is None:
            spec.errors.append(f'grid row {r}: both from_month and max_import_mw are needed'); continue
        grid.append((int(round(a)), float(b)))
    if grid:
        spec.grid_schedule = grid
    validate(spec)
    return spec


def validate(spec):
    e, w = spec.errors, spec.warnings
    chk = lambda k, lo, hi=None: getattr(spec, k) is not None and (getattr(spec, k) < lo or (hi is not None and getattr(spec, k) > hi))
    for k, lo in [('onsite_energy_land_acres', 0), ('gas_engine_land_acres_per_mw', 0), ('fuel_cell_land_acres_per_mw', 0), ('relocation_notice_minutes', 0),
                  ('interruption_consequence_usd_per_mwh', 0), ('remote_capacity_mw', 0), ('remote_electricity_usd_per_mwh', 0),
                  ('relocation_service_usd_per_local_mwh', 0), ('nuclear_payment_usd_per_mwh', 0), ('nuclear_max_mw', 0), ('nuclear_available_month', 0)]:
        if chk(k, lo):
            e.append(f'{k} cannot be negative')
    if chk('remote_energy_multiplier', 1.0):
        e.append('remote_energy_multiplier must be at least 1.00 (1.00 = no overhead)')
    if spec.mature_peak_demand_mw is not None and spec.mature_peak_demand_mw <= 0:
        e.append('mature_peak_demand_mw must be positive')
    if spec.default_initial_service_month is not None and spec.default_initial_service_month < int(cs.case['campus_ready_month']):
        e.append(f"default_initial_service_month cannot be earlier than the earliest physical readiness (month {int(cs.case['campus_ready_month'])})")
    if spec.grid_delay_months is not None and spec.grid_delay_months < 0:
        e.append('grid_delay_months cannot be negative')
    if spec.nuclear_max_mw is not None and spec.nuclear_max_mw % NUCLEAR_BLOCK > 1e-9:
        fl = float(int(spec.nuclear_max_mw // NUCLEAR_BLOCK) * NUCLEAR_BLOCK)
        w.append(f'nuclear_max_mw {spec.nuclear_max_mw:g} is not a multiple of the {NUCLEAR_BLOCK:g} MW block; using {fl:g}')
        spec.nuclear_max_mw = fl
    if spec.ramp:
        months = [m for m, _ in spec.ramp]
        if months != list(range(len(months))):
            e.append('Load ramp must list months 0, 1, 2, ... in order')
        if any(f < 0 or f > 1 for _, f in spec.ramp):
            e.append('Load fractions must be between 0 and 1')
        if abs(spec.ramp[-1][1] - 1.0) > 1e-9:
            e.append('Load ramp must end at 100% (1.0): the model holds demand flat after the last row, so the last row has to be full demand')
        if any(b < a - 1e-12 for (_, a), (_, b) in zip(spec.ramp, spec.ramp[1:])):
            w.append('Load ramp decreases between some months; the model uses it as given')
        if spec.ramp_months() > 40:
            w.append('Ramp longer than 40 months: pools and deadline logic were tested with the 10-month packet ramp')
    if spec.grid_schedule:
        g = spec.grid_schedule
        if [m for m, _ in g] != sorted({m for m, _ in g}):
            e.append('Grid schedule months must be unique and in ascending order')
        if g and g[0][0] != 0:
            w.append('Grid schedule does not start at month 0; assuming 0 MW before the first step')
            spec.grid_schedule = [(0, 0.0)] + g
        if any(x < 0 for _, x in g):
            e.append('Grid import capacity cannot be negative')
    if spec.mature_peak_demand_mw is not None and spec.grid_schedule:
        top = max(x for _, x in spec.grid_schedule)
        if top < spec.mature_peak_demand_mw * 0.6:
            w.append(f'Largest grid step ({top:g} MW) is below 60% of mature peak demand: expect heavy reliance on onsite generation')
    return spec


# ----------------------------------------------------------------------------------------------- applying a case (per thread, nothing global is mutated)
import casepatch as CPH  # noqa: E402  (installs the per-thread proxies on first import)

_REG, _REG_LOCK = {}, threading.Lock()


def _build_tables(spec, key):
    """Immutable copies of the packet tables carrying this case's values. Built once per case."""
    base = CPH.BENCH
    ov = spec.overrides()
    opts = base.options.copy()
    row = lambda oid: opts.option_id == oid
    if 'gas_engine_land_acres_per_mw' in ov: opts.loc[row('gas_engine'), 'land_acres_per_mw'] = ov['gas_engine_land_acres_per_mw']
    if 'fuel_cell_land_acres_per_mw' in ov: opts.loc[row('fuel_cell'), 'land_acres_per_mw'] = ov['fuel_cell_land_acres_per_mw']
    if 'nuclear_available_month' in ov: opts.loc[row('nuclear'), 'available_month'] = int(ov['nuclear_available_month'])
    if 'nuclear_payment_usd_per_mwh' in ov: opts.loc[row('nuclear'), 'payment_usd_per_mwh'] = ov['nuclear_payment_usd_per_mwh']
    if 'nuclear_max_mw' in ov: opts.loc[row('nuclear'), 'max_mw'] = ov['nuclear_max_mw']
    case = copy.deepcopy(base.case)
    if 'onsite_energy_land_acres' in ov: case['onsite_energy_land_acres'] = ov['onsite_energy_land_acres']
    resp = dict(base.response)
    for k in ('relocation_notice_minutes', 'interruption_consequence_usd_per_mwh', 'remote_capacity_mw', 'remote_energy_multiplier',
              'remote_electricity_usd_per_mwh', 'relocation_service_usd_per_local_mwh'):
        if k in ov: resp[k] = ov[k]
    ramp = base.ramp
    if 'ramp' in ov:
        ramp = pd.DataFrame({'months_since_initial_service': [m for m, _ in ov['ramp']], 'load_fraction': [f for _, f in ov['ramp']]})
        ramp['equivalent_peak_mw'] = ramp.load_fraction * spec._val('mature_peak_demand_mw')
        case['ramp_duration_months'] = spec.ramp_months()
    grid = base.grid_schedule
    if 'grid_schedule' in ov:
        grid = pd.DataFrame(ov['grid_schedule'], columns=['from_month', 'max_import_mw'])
        case['grid_access'] = [{'from_month': int(m), 'max_import_mw': float(x)} for m, x in ov['grid_schedule']]
    hourly = base.hourly
    if 'mature_peak_demand_mw' in ov:
        f = ov['mature_peak_demand_mw'] / REFERENCE['mature_peak_demand_mw']
        hourly = hourly.copy()
        for c in ('load_mw', 'tier0_protected_mw', 'tier1_relocatable_mw', 'tier2_interruptible_mw'):
            hourly[c] = hourly[c] * f
        case['peak_mw'] = ov['mature_peak_demand_mw']
    if 'stress_events' in ov:
        hourly = hourly.copy()
        ts = pd.to_datetime(hourly['timestamp_utc'], utc=True)
        active = np.zeros(len(hourly), bool)
        frac = np.ones(len(hourly))
        notice = np.zeros(len(hourly))
        ids = np.array(['none'] * len(hourly), dtype=object)
        for e in ov['stress_events']:
            m = ((ts >= pd.Timestamp(e['start_utc'])) & (ts < pd.Timestamp(e['end_utc']))).to_numpy()
            active |= m; frac[m] = float(e['import_fraction']); notice[m] = float(e.get('notice_minutes', 30)); ids[m] = str(e['id'])
        hourly['stress_active'] = active.astype(int); hourly['stress_import_fraction'] = frac; hourly['notice_minutes'] = notice.astype(int); hourly['stress_event_id'] = ids
    o = opts.copy(); o['option_id'] = o['option_id'].astype(str)
    e_opt = opts.set_index('option_id')
    return CPH.Tables(key, opts, case, resp, ramp, grid, hourly, o.set_index('option_id', drop=False), e_opt,
                      {k: E.NUM(e_opt.loc[k, 'max_mw']) for k in e_opt.index}, {k: E.NUM(e_opt.loc[k, 'block_mw']) for k in e_opt.index})


def tables_for(spec):
    key = 'bench' if spec is None or spec.is_default() else spec.key()
    if key == 'bench':
        return CPH.BENCH
    with _REG_LOCK:
        if key not in _REG:
            _REG[key] = _build_tables(spec, key)
        return _REG[key]


def set_case(spec):
    """Select this case for the CURRENT THREAD (None or a default spec = the packet benchmark). Other threads are unaffected."""
    t = tables_for(spec)
    CPH.set_tables(t)
    return t.key


activate = set_case          # older name


@contextmanager
def using(spec):
    """Scoped selection for the current thread."""
    with CPH.using_tables(tables_for(spec)):
        yield


def active_key():
    return CPH.current().key
