"""Compare GridTech_Inputs.xlsx (a formula-based planning/screening workbook) with our model.

    python team_app/compare_excel.py /path/to/GridTech_Inputs.xlsx

1  facts the workbook states (computed from the packet, plus its formulas) vs the same quantities from our evaluator
2  its four sample portfolios (placeholders, not model output) vs our evaluator on reconstructed portfolios and vs our best at the same opening month,
   and what its "onsite backup >= Tier 0 peak" rule costs
3  its eight customer types (weights + hard limits) run through our filters and scoring, vs the pick its workbook formulas make
Writes outputs/excel_comparison/*.csv and COMPARISON.md."""
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import alt_adapter as A  # noqa: E402
import engine as E  # noqa: E402
import logic as L  # noqa: E402
import risk as R  # noqa: E402

gm, cs = A.gm, E.cs
OUT = HERE.parent / 'outputs' / 'excel_comparison'


def close(a, b, rel=0.005, abs_=0.05):
    return abs(a - b) <= max(abs_, rel * abs(b))


# ------------------------------------------------------------------------------------------------ 1. facts
def facts(wb):
    kn, rd = wb['Key numbers'], wb['Resource data']
    ctx = gm.get_ctx(20, 0)
    fr = ctx.base['frame']
    rows = []
    add = lambda area, item, theirs, ours, unit='': rows.append(dict(area=area, item=item, workbook=theirs, ours=ours, unit=unit, status='match' if close(float(theirs), float(ours)) else 'DIFFERENT'))
    ordinary = fr.ordinary_grid_limit_mw.to_numpy()
    for r in range(6, 13):
        m = int(kn.cell(r, 1).value)
        sel = ctx.month == m
        add('Grid gap', f'max (load - ordinary grid cap), month {m}', kn.cell(r, 2).value, float((ctx.load[sel] - ordinary[sel]).max()), 'MW')
    add('Grid gap', 'peak gap before the 75 MW connection', kn['B13'].value, max(float((ctx.load[ctx.month == m] - ordinary[ctx.month == m]).max()) for m in range(21, 30)), 'MW')
    full = ctx.month >= 30
    act = ctx.active.astype(bool)
    add('Stress', 'stress hours per model year', kn['B17'].value, int(act[ctx.year == 0].sum()), 'h')
    s = full & act
    add('Stress', 'max shortfall vs 60% cap, no flexibility', kn['B18'].value, float((ctx.load[s] - ctx.grid_limit[s]).max()), 'MW')
    add('Stress', 'max shortfall after relocation + interruption', kn['B19'].value, float((ctx.load[s] - np.minimum(ctx.tier1[s], 12 / 1.05) - ctx.tier2[s] - ctx.grid_limit[s]).max()), 'MW')
    add('Stress', 'max Tier 0 above the 60% cap (hard floor of onsite supply)', kn['B20'].value, float((ctx.tier0[s] - ctx.grid_limit[s]).max()), 'MW')
    add('Load', 'Tier 0 (protected) peak, full demand', kn['B58'].value, float(ctx.tier0.max()), 'MW')
    add('Load', 'Tier 0 peak during stress hours', kn['B59'].value, float(ctx.tier0[act].max()), 'MW')
    add('Flexibility', 'relocation cost per local MWh', kn['B27'].value, float(ctx.reloc_price), '$/MWh')
    add('Flexibility', 'max relocation (local-equivalent)', kn['B29'].value, float(gm.RESP['remote_capacity_mw'] / gm.RESP['remote_energy_multiplier']), 'MW')
    add('Flexibility', 'interruption / relocation cost ratio', kn['B31'].value, float(ctx.intr_price / ctx.reloc_price), 'x')
    for r in range(36, 46):
        y = int(kn.cell(r, 1).value)
        sel = ctx.year == y - 1
        add('Prices', f'model year {y}: merchant $/MWh (incl. adder)', kn.cell(r, 2).value, float(ctx.grid_price[sel].mean()), '$/MWh')
        add('Prices', f'model year {y}: delivered gas $/MMBtu', kn.cell(r, 3).value, float(ctx.gas_price[sel].mean()), '$/MMBtu')
        add('Prices', f'model year {y}: grid CO2 t/MWh', kn.cell(r, 4).value, float(ctx.grid_co2[sel].mean()), 't/MWh')
    nuc = gm.OPT.loc['nuclear']
    add('Nuclear', 'annual payment (pay-for-available)', kn['B51'].value, float(nuc.max_mw * 8760 * nuc.payment_usd_per_mwh), '$/yr')
    add('Nuclear', 'share of mature annual load', kn['B52'].value, float(nuc.max_mw * 8760 / gm.CASE['full_load_year_mwh']), 'fraction')
    bat = gm.OPT.loc['battery_4h']
    deliver = bat.max_mw * bat.duration_hours * np.sqrt(bat.round_trip_efficiency)
    add('Battery', 'deliverable energy of the 12 MW battery', kn['B62'].value, float(deliver), 'MWh')
    add('Battery', 'ride-through at the Tier 0 peak', kn['B63'].value, float(deliver / ctx.tier0.max() * 60), 'min')
    for r in range(11, 18):
        oid = rd.cell(r, 2).value
        o = gm.OPT.loc[oid]
        crf = 0.0 if not o.capex_usd_per_kw else gm.CASE['discount_rate_real'] / (1 - (1 + gm.CASE['discount_rate_real']) ** (-o.life_years))
        add('Resources', f'{oid}: CRF', rd.cell(r, 16).value, float(crf), '')
        add('Resources', f'{oid}: annual fixed $/kW-yr (CAPEX x CRF + fixed O&M)', rd.cell(r, 17).value, float(o.capex_usd_per_kw * crf + o.fom_usd_per_kw_year), '$/kW-yr')
        add('Resources', f'{oid}: CO2 kg/MWh', rd.cell(r, 19).value, float(o.co2_t_per_mwh * 1000), 'kg/MWh')
        mx = o.max_mw if not o.land_acres_per_mw else min(o.max_mw, int(gm.CASE['onsite_energy_land_acres'] / o.land_acres_per_mw / o.block_mw) * o.block_mw)
        add('Resources', f'{oid}: max MW within the 8 acres', rd.cell(r, 20).value, float(mx), 'MW')
    add('Resources', 'avg solar capacity factor', rd['O11'].value, float(cs.hourly.solar_proxy_cf.mean()), '')
    add('Resources', 'avg wind capacity factor', rd['O12'].value, float(cs.hourly.wind_proxy_cf.mean()), '')
    add('Resources', 'avg delivered gas price over 87,600 h', rd['B7'].value, float(ctx.gas_price.mean()), '$/MMBtu')
    # sample-portfolio fixed maintenance
    fom = {k: float(gm.OPT.loc[k].fom_usd_per_kw_year) for k in ('gas_turbine', 'gas_engine', 'fuel_cell', 'battery_4h')}
    for col, name in zip('CDEF', ['Balanced', 'Speed-first', 'Cost-first', 'Carbon-first']):
        mw = [kn[f'{col}{r}'].value for r in range(67, 71)]
        add('Maintenance', f'{name} sample: annual fixed O&M', kn[f'{col}73'].value, float(sum(m * 1000 * f for m, f in zip(mw, fom.values()))), '$/yr')
    return pd.DataFrame(rows)


# ------------------------------------------------------------------------------------------------ 2. sample portfolios
def sample_portfolios(wb):
    sd, kn = wb['Scoring demo'], wb['Key numbers']
    cols = {'A': 'D', 'B': 'C', 'C': 'E', 'D': 'F'}                          # candidate -> Key numbers column of onsite MW
    onsite = {c: dict(gas_turbine=kn[f'{k}67'].value, gas_engine=kn[f'{k}68'].value, fuel_cell=kn[f'{k}69'].value, battery_4h=kn[f'{k}70'].value) for c, k in cols.items()}
    # contracts are described only in words ("60 MW wind", "oversized wind + solar, 25 MW nuclear"): our stated reading
    contracts = {'A': {}, 'B': dict(wind=60.0), 'C': dict(wind=60.0), 'D': dict(wind=150.0, utility_solar=100.0, nuclear=25.0)}
    pool = L.load_pool('base')
    rows, pf = [], {}
    for i, r in enumerate(range(6, 10)):
        label, start = sd.cell(r, 1).value, int(sd.cell(r, 2).value)
        c = 'ABCD'[i]
        caps = {k: float(v) for k, v in {**onsite[c], **contracts[c]}.items() if v}
        m, res, ctx = gm.evaluate(start, caps, A.scenario())
        g = pool[pool.start == start]
        cheapest, cleanest = g.loc[g.total_cost.idxmin()], g.loc[g.intensity.idxmin()]
        rows.append({'candidate': label, 'opening': start, 'workbook cost $M': sd.cell(r, 4).value, 'workbook kg/MWh': sd.cell(r, 5).value,
                     'workbook maintenance $M': sd.cell(r, 10).value, 'reconstructed portfolio': ', '.join(f'{k} {v:g}' for k, v in caps.items()),
                     'our cost $M': round(m['total_decision_cost_usd'] / 1e6, 1), 'our kg/MWh': round(m['combined_electricity_co2_kg_per_mwh']), 'feasible': bool(res['feasible']),
                     'our maintenance $M': round(R.maintenance(caps, start, A.simulate(caps, start)['_hourly'])['total'] / 1e6, 1) if res['feasible'] else float('nan'),
                     'our cheapest at this month $M': round(cheapest.total_cost / 1e6, 1), 'its kg/MWh': round(cheapest.intensity),
                     'our cleanest at this month kg/MWh': round(cleanest.intensity), 'its cost $M': round(cleanest.total_cost / 1e6, 1)})
    return pd.DataFrame(rows)


def reliability_floor_cost():
    """Price of the workbook's rule 'onsite generators + battery >= Tier 0 peak (51.1 MW)', and of a stricter N-1 reading, from our pool."""
    pool = L.load_pool('base')
    crit = float(gm.get_ctx(20, 0).tier0.max())
    tot = pool.gas_turbine + pool.gas_engine + pool.fuel_cell + pool.battery_4h
    largest = np.where(pool.gas_turbine > 0, 25.0, np.where((pool.gas_engine + pool.fuel_cell) > 0, 5.0, 0.0))
    firm = (pool.gas_turbine + pool.gas_engine + pool.fuel_cell) - largest
    rows = []
    for s in (14, 20, 24, 26):
        g = pool.start == s
        base = pool[g].total_cost.min()
        w = pool[g & (tot >= crit)]
        n1 = pool[g & (firm >= crit)]
        rows.append({'opening': s, 'cheapest, no reliability rule $M': round(base / 1e6, 1),
                     'cheapest with workbook rule (gas+battery >= 51.1 MW) $M': round(w.total_cost.min() / 1e6, 1) if len(w) else float('nan'),
                     'extra $M': round((w.total_cost.min() - base) / 1e6, 1) if len(w) else float('nan'),
                     'cheapest with N-1 firm gas >= 51.1 MW $M': round(n1.total_cost.min() / 1e6, 1) if len(n1) else float('nan'),
                     'extra (N-1) $M': round((n1.total_cost.min() - base) / 1e6, 1) if len(n1) else float('nan')})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------------------------------------ 3. customer types
def customer_types(wb):
    ct, sd = wb['Customer types'], wb['Scoring demo']
    pools = {'base': L.load_pool('base'), 'no_t2': L.load_pool('no_t2'), 'd6': L.load_pool('d6')}
    cand = {sd.cell(13, c).value: sd.cell(r, 1).value for c, r in zip(range(8, 12), range(6, 10))}      # 'A' -> label
    cand_cost = {sd.cell(13, c).value: (sd.cell(r, 4).value, sd.cell(r, 5).value) for c, r in zip(range(8, 12), range(6, 10))}
    rows = []
    for i, r in enumerate(range(6, 14)):
        name = ct.cell(r, 1).value
        full_by, max_co2, prem = ct.cell(r, 4).value, ct.cell(r, 5).value, ct.cell(r, 6).value
        t2_ok, delay = ct.cell(r, 7).value, int(ct.cell(r, 8).value or 0)
        w = (ct.cell(r, 9).value, ct.cell(r, 10).value, ct.cell(r, 11).value)
        value = (ct.cell(r, 13).value or 0) / 1e6
        months = [int(x) for x in re.findall(r'\d+', str(ct.cell(r, 14).value))] if ct.cell(r, 14).value else []
        their_pick = sd.cell(14 + i, 12).value
        if name.startswith('Staged'):
            rows.append({'customer type': name, 'workbook pick': 'B (same as balanced)', 'our pick': 'not modelled (staged capacity, optional extension)'}); continue
        pname = 'd6' if delay == 6 else ('base' if str(t2_ok).lower().startswith('y') else 'no_t2')
        pool = pools[pname]
        base_cost = pools['base'][pools['base'].start == 20].total_cost.min()                       # month-30 reference for the premium test
        req = {}
        if isinstance(full_by, (int, float)): req['deadline_full'] = full_by
        if isinstance(max_co2, (int, float)): req['carbon_limit'] = max_co2
        if isinstance(prem, (int, float)): req['premium_pct'] = prem * 100
        surv, filt = L.apply_requirements(pool, req, (), base_cost if 'premium_pct' in req else None)
        row = {'customer type': name, 'weights S/C/E': '/'.join(str(int(x)) for x in w), 'hard limits': ', '.join(f'{k}={v}' for k, v in req.items()) or 'none',
               'pool used': {'base': 'base', 'no_t2': 'no Tier 2', 'd6': '6-month grid delay'}[pname],
               'workbook pick': (f"{their_pick}: {cand.get(their_pick, '')[:34]} (placeholder ${cand_cost[their_pick][0]:.0f}M, {cand_cost[their_pick][1]:.0f} kg)" if their_pick in cand else their_pick)}
        if surv.empty:
            from_rows = L.explain_empty(pool, filt, ())
            row['our pick'] = 'no portfolio passes: ' + (f"{from_rows[0][1]} binds; needs {from_rows[0][3]:,.0f} ({from_rows[0][2]})" if from_rows else 'several limits conflict')
        else:
            best = L.rank(surv, w, value, 0.0, ref=pool).iloc[0]
            row.update({'our pick': f"open {int(best.start)} -> full {int(best.full_month)}, ${best.total_cost / 1e6:.0f}M, {best.intensity:.0f} kg/MWh: {L.caps_text(best)}",
                        'our opening': int(best.start), 'our cost $M': round(best.total_cost / 1e6, 1), 'our kg/MWh': round(best.intensity), 'passing portfolios': len(surv)})
            if months:
                s2 = surv[surv.start.isin(months)]
                if len(s2):
                    b2 = L.rank(s2, w, value, 0.0, ref=pool).iloc[0]
                    row['our pick in the months the workbook tests'] = f"open {int(b2.start)}, ${b2.total_cost / 1e6:.0f}M, {b2.intensity:.0f} kg/MWh"
                else:
                    row['our pick in the months the workbook tests'] = f'nothing passes in months {months}'
        row['workbook expected direction'] = ct.cell(r, 15).value
        rows.append(row)
    return pd.DataFrame(rows)


def main(path):
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(path, data_only=True)
    f = facts(wb); f.to_csv(OUT / '1_facts.csv', index=False)
    s = sample_portfolios(wb); s.to_csv(OUT / '2_sample_portfolios.csv', index=False)
    rel = reliability_floor_cost(); rel.to_csv(OUT / '2b_reliability_floor.csv', index=False)
    c = customer_types(wb); c.to_csv(OUT / '3_customer_types.csv', index=False)
    pd.set_option('display.width', 250, 'display.max_colwidth', 120, 'display.max_columns', 30)
    print(f"1. FACTS: {len(f)} compared, {int((f.status == 'match').sum())} match, {int((f.status != 'match').sum())} different")
    print(f[f.status != 'match'].to_string(index=False) if (f.status != 'match').any() else '   all match')
    print('\n2. SAMPLE PORTFOLIOS\n' + s.drop(columns=['reconstructed portfolio']).to_string(index=False))
    print(s[['candidate', 'reconstructed portfolio']].to_string(index=False))
    print('\n2b. PRICE OF THE WORKBOOK RELIABILITY RULE\n' + rel.to_string(index=False))
    print('\n3. CUSTOMER TYPES')
    for _, r in c.iterrows():
        print(f"- {r['customer type']} | weights {r.get('weights S/C/E', '')} | limits {r.get('hard limits', '')}\n    workbook: {r['workbook pick']}\n    ours:     {r['our pick']}"
              + (f"\n    ours (workbook's months only): {r['our pick in the months the workbook tests']}" if isinstance(r.get('our pick in the months the workbook tests'), str) else ''))
    return f, s, rel, c


if __name__ == '__main__':
    main(sys.argv[1])
